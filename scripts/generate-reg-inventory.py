#!/usr/bin/env python3
"""Regenerate lib/regInventory.ts from the OJK regulation inventory workbook.

Usage:
    python scripts/generate-reg-inventory.py "/path/to/Inventory Reg OJK_<Month>_<Year>_Regsmap site.xlsx"

Reads the "Reg_Master_CLEAN" sheet (row 2 = header, data from row 3). The ID
cell's hyperlink decides how the row links out, exactly like the workbook's
legend says: a betterbankings.com/regmaps/... target is an internal RegMaps
link (orange in the sheet), anything else is the official OJK document (blue).

The "DSU Note" column is internal editorial tracking and is deliberately not
published.
"""

import json
import sys
from datetime import datetime

import openpyxl

SHEET = "Reg_Master_CLEAN"
HEADER_ROW = 2
INTERNAL_HOST = "betterbankings.com"


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split()).strip()


def format_date(value):
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y")
    return clean(value)


def split_refs(value):
    """Reference columns list one entry per line; one row uses " | " instead."""
    text = str(value or "").replace(" | ", "\n")
    parts = [" ".join(p.split()).strip() for p in text.split("\n")]
    return [p for p in parts if p]


# The sheet's Category column has inconsistent casing and trailing spaces.
CATEGORIES = {
    "business and organization": "Business and Organization",
    "capital & icaap": "Capital & ICAAP",
    "credit risk": "Credit Risk",
    "disclosure and reporting": "Disclosure and Reporting",
    "financial conglomeration": "Financial Conglomeration",
    "governance and risk management": "Governance and Risk Management",
    "irrbb": "IRRBB",
    "it and cybersecurity": "IT and Cybersecurity",
    "licensing": "Licensing",
    "liquidity risk and ilaap": "Liquidity Risk and ILAAP",
    "market, counterparty and cva risk": "Market, Counterparty and CVA Risk",
    "operational risk & resilience": "Operational Risk & Resilience",
}


def normalize_category(value):
    text = clean(value)
    return CATEGORIES.get(text.lower(), text)


def main(path, out):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    header = [clean(c.value) for c in ws[HEADER_ROW]]
    idx = {name: i for i, name in enumerate(header) if name}

    entries = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        id_cell = row[idx["ID"]]
        if not clean(id_cell.value):
            continue

        target = id_cell.hyperlink.target if id_cell.hyperlink else None
        if target and INTERNAL_HOST in target:
            # Strip the origin so the link stays inside the Next.js router.
            link = target[target.index("/regmaps") :]
            link_type = "regmaps"
        elif target:
            link = target
            link_type = "ojk"
        else:
            link = ""
            link_type = "none"

        entries.append(
            {
                "status": clean(row[idx["Status"]].value),
                "id": clean(id_cell.value),
                "type": clean(row[idx["Type"]].value),
                # A couple of rows use "BUK. BUS, UUS" — normalize the typo.
                "applicability": clean(row[idx["Applicability"]].value).replace(". ", ", "),
                "date": format_date(row[idx["Date"]].value),
                "title": clean(row[idx["Title"]].value),
                "category": normalize_category(row[idx["Category"]].value),
                "code": clean(row[idx["Code"]].value),
                "revokes": split_refs(row[idx["Revokes"]].value),
                "amendedBy": split_refs(row[idx["Amended By"]].value),
                "replacedBy": split_refs(row[idx["Replaced By"]].value),
                "link": link,
                "linkType": link_type,
            }
        )

    body = ",\n".join("  " + json.dumps(e, ensure_ascii=False) for e in entries)
    source = path.replace("\\", "/").rsplit("/", 1)[-1]
    with open(out, "w", encoding="utf-8") as f:
        f.write(
            f"""// AUTO-GENERATED — do not edit by hand.
// Source: "{source}", sheet "{SHEET}".
// Regenerate with: python scripts/generate-reg-inventory.py "<workbook>.xlsx"
// The workbook's "DSU Note" column is internal tracking and is not published.

export type RegLinkType = "regmaps" | "ojk" | "none";

export interface RegInventoryEntry {{
  status: string;
  id: string;
  type: string;
  applicability: string;
  date: string;
  title: string;
  category: string;
  code: string;
  revokes: string[];
  amendedBy: string[];
  replacedBy: string[];
  /** Internal /regmaps path when linkType is "regmaps", absolute OJK URL otherwise. */
  link: string;
  linkType: RegLinkType;
}}

export const regInventory: RegInventoryEntry[] = [
{body},
];
"""
        )
    print(f"wrote {len(entries)} entries to {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "lib/regInventory.ts")
